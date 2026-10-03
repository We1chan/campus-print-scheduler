"""Generate real two-page Office inputs for preview integration acceptance.

Test-only dependencies: python-docx, python-pptx, openpyxl. LibreOffice converts
the OOXML inputs to their legacy counterparts in an isolated temporary profile.
No user documents or global Office profile are opened or modified.
"""
import argparse
import os
from pathlib import Path
import subprocess
import tempfile


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--office", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    root = args.output_dir.resolve()
    if root.exists():
        parser.error("--output-dir must not exist")
    root.mkdir(parents=True)
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.pagebreak import Break

    title = "中文 空格"
    doc = Document()
    doc.add_heading("Paperroom Word page ONE", 0)
    doc.add_paragraph("Actual formatted document content: 第一页，保留分页。")
    doc.add_page_break()
    doc.add_heading("Paperroom Word page TWO", 0)
    doc.add_paragraph("Second page: 第二页，不是文件信息占位。")
    doc.save(root / (title + ".docx"))
    ppt = Presentation()
    ppt.slide_width = Inches(13.333333)
    ppt.slide_height = Inches(7.5)
    for number in (1, 2):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        box.text = f"Paperroom slide {number}\n真实幻灯片内容 {number}"
    ppt.save(root / (title + ".pptx"))
    book = Workbook()
    sheet = book.active
    sheet.title = "打印区域"
    for row in range(1, 51):
        sheet.cell(row, 1, f"Paperroom row {row}")
        sheet.cell(row, 2, row * 10)
        sheet.cell(row, 1).fill = PatternFill("solid", fgColor="FFE0D0" if row <= 25 else "D0E0FF")
        sheet.row_dimensions[row].height = 18
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 15
    sheet["Z100"] = "OUTSIDE PRINT AREA - must not create another page"
    sheet.print_area = "A1:B50"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.scale = 100
    sheet.row_breaks.append(Break(id=25))
    book.save(root / (title + ".xlsx"))
    with tempfile.TemporaryDirectory(prefix="paperroom-fixture-profile-") as directory:
        for source, destination, filter_name in [("docx", "doc", "MS Word 97"),
                ("pptx", "ppt", "MS PowerPoint 97"), ("xlsx", "xls", "MS Excel 97")]:
            subprocess.run([str(args.office.resolve()), "-env:UserInstallation=" + Path(directory).as_uri(),
                "--headless", "--nologo", "--norestore", "--convert-to", destination + ":" + filter_name,
                "--outdir", str(root), str(root / (title + "." + source))],
                env=dict(os.environ, PYTHONDONTWRITEBYTECODE="1"), check=True, timeout=90)
            if not (root / (title + "." + destination)).is_file():
                raise RuntimeError("Legacy fixture conversion did not produce " + destination)
    print(root)


if __name__ == "__main__":
    main()
